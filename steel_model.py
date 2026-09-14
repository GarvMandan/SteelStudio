"""JSON and geometry adapter for the existing structural steel takeoff engine.

Coordinates are feet, with Z up. Quantities and demands come from
calculation_engine; rendered profiles are schematic, not connection details.
Section selections are takeoff inputs and never an engineering approval.
"""
from __future__ import annotations

import copy
import math
import re
from functools import lru_cache
from pathlib import Path

import calculation_engine as engine
from calculation_engine import InputValidationError, axis_index, axis_letter


MAX_BAYS_PER_AXIS = 24
MAX_JOIST_SPACES = 40
MAX_MEMBERS = 18000
DESCRIPTIONS = {
    "joist": "Open-web steel joist. Parallel chords and diagonal web members span between supporting girders or walls. Catalog depth is nominal; the web pattern shown is schematic.",
    "girder": "Joist girder. An open-web primary member supporting joist reactions at panel points. Depth, panel count and panel load define the catalog entry; its weight also depends on span.",
    "column": "Hollow structural section (HSS) column. A closed steel tube supporting the framing above. Nominal outside dimensions describe the section; steel grade and connection design are unspecified.",
}


def defaults():
    """Return an independent, editable example project with no sections assigned."""
    return {
        "app": "Structural Steel Studio", "visualizer_version": 2,
        "name": "New steel structure", "x_spans_ft": [40.0, 40.0, 40.0],
        "y_spans_ft": [40.0, 40.0], "joist_spaces": 6,
        "roof_height_ft": 24.0, "roof_type": "Single Slope", "roof_rise_ft": 0.0,
        "clear_height_ft": 24.0, "roof_calculation_mode": "original",
        "auto_select_sections": True, "break_clear_height": True,
        "speed_bay_rows": [], "footing_depth_ft": 1.0, "bearing_capacity_psf": 3000.0,
        "metal_deck_thickness_in": 1.5, "insulation_depth_in": 3.0,
        "single_slope_direction": "North", "joist_seat_depth_in": 2.5,
        "load_inputs_psf": {"dead_load_psf": 20.0, "live_load_psf": 20.0,
                            "snow_load_psf": 30.0, "snow_code": "ASCE 7-16",
                            "reduced_snow_load_psf_manual": 0.0,
                            "collateral_addition_psf": 5.0},
        "inactive_bays": [], "collateral_bays": [], "additional_load_layers": [],
        "joists_per_bay": [], "load_bearing_wall_segments": [],
        "mezzanine_enabled": False, "mezzanines": [], "assignments": {},
        "member_overrides": {}, "group_overrides": {}, "default_sections": {},
    }


def _number(value, label, minimum=0.0, maximum=100000.0, integer=False):
    try:
        number = float(value)
    except (ValueError, TypeError, OverflowError) as exc:
        raise InputValidationError(f"{label} must be a number.") from exc
    if isinstance(value, bool) or not math.isfinite(number) or not minimum <= number <= maximum:
        raise InputValidationError(f"{label} must be finite and between {minimum:g} and {maximum:g}.")
    if integer and number != int(number):
        raise InputValidationError(f"{label} must be a whole number.")
    return int(number) if integer else number


def _finite_tree(value, path="project", depth=0):
    if depth > 30:
        raise InputValidationError("Project nesting is too deep.")
    if isinstance(value, float) and not math.isfinite(value):
        raise InputValidationError(f"{path} contains a non-finite number.")
    if isinstance(value, dict):
        for key, child in value.items():
            _finite_tree(child, f"{path}.{key}", depth + 1)
    elif isinstance(value, list):
        if len(value) > 30000:
            raise InputValidationError(f"{path} contains too many entries.")
        for i, child in enumerate(value):
            _finite_tree(child, f"{path}[{i}]", depth + 1)


def _list(value, label, maximum=1500):
    if not isinstance(value, list) or len(value) > maximum:
        raise InputValidationError(f"{label} must be an array of at most {maximum} entries.")
    return value


def _bay_indices(items, nx, ny, label):
    result = set()
    for item in _list(items or [], label):
        if not isinstance(item, dict):
            raise InputValidationError(f"{label}: each bay needs x_bay and y_bay.")
        x = _number(item.get("x_bay"), f"{label} X bay", 1, MAX_BAYS_PER_AXIS, True) - 1
        y = axis_index(item.get("y_bay", ""))
        # Resizing the grid prunes old bay-specific edits outside its new bounds.
        if x < nx and 0 <= y < ny:
            result.add((x, y))
    return result


def _serialize_bays(bays):
    return [{"x_bay": x + 1, "y_bay": axis_letter(y)} for x, y in sorted(bays)]


def normalize_project(project):
    if not isinstance(project, dict):
        raise InputValidationError("Project must be a JSON object.")
    _finite_tree(project)
    p = copy.deepcopy(project)
    base = defaults()
    legacy = isinstance(p.get("grid"), dict) and "x_spans_ft" not in p
    grid = p.get("grid") or {}
    inputs = p.get("inputs") or {}
    if not isinstance(grid, dict) or not isinstance(inputs, dict):
        raise InputValidationError("Project grid and inputs must be objects.")
    if legacy:
        for key, value in grid.items():
            p.setdefault(key, copy.deepcopy(value))
        for key in ("mezzanine_enabled", "single_slope_direction", "joist_seat_depth_in"):
            if key in inputs:
                p.setdefault(key, inputs[key])
        p.setdefault("joist_spaces", inputs.get("joist_spaces_default", base["joist_spaces"]))
        p.setdefault("roof_height_ft", inputs.get("clear_height_ft", base["roof_height_ft"]))
        p.setdefault("roof_type", inputs.get("roof_type", "Flat"))
        p.setdefault("load_inputs_psf", {})
        for key in base["load_inputs_psf"]:
            if key in inputs:
                p["load_inputs_psf"].setdefault(key, inputs[key])
        if "reduced_snow_load_manual_psf" in inputs:
            p["load_inputs_psf"].setdefault("reduced_snow_load_psf_manual", inputs["reduced_snow_load_manual_psf"])
        p["legacy_roof_profile"] = True
    # Preserve the first Studio's explicit eave/rise geometry on migration.
    p.setdefault("roof_calculation_mode", "original" if legacy else ("eave" if project.get("visualizer_version") == 1 else "original"))
    p.setdefault("clear_height_ft", inputs.get("clear_height_ft", max(4, float(p.get("roof_height_ft", 24)) - (4 if p["roof_calculation_mode"] == "eave" else 0))))
    for key in ("footing_depth_ft", "metal_deck_thickness_in", "insulation_depth_in", "break_clear_height"):
        if key in inputs:
            p.setdefault(key, inputs[key])
    p.setdefault("bearing_capacity_psf", inputs.get("bearing_pressure_psi", 3000))
    for key, value in base.items():
        p.setdefault(key, copy.deepcopy(value))
    for key in ("x_spans_ft", "y_spans_ft"):
        spans = _list(p[key], key, MAX_BAYS_PER_AXIS)
        if not spans:
            raise InputValidationError("Add at least one bay on each axis.")
        p[key] = [_number(v, key, .5, 200) for v in spans]
    nx, ny = len(p["x_spans_ft"]), len(p["y_spans_ft"])
    p["name"] = str(p.get("name") or "Untitled structure")[:200]
    p["joist_spaces"] = _number(p["joist_spaces"], "Joist spaces", 1, MAX_JOIST_SPACES, True)
    p["roof_height_ft"] = _number(p["roof_height_ft"], "Roof elevation", 4, 200)
    p["clear_height_ft"] = _number(p["clear_height_ft"], "Clear height", 4, 200)
    p["footing_depth_ft"] = _number(p["footing_depth_ft"], "Footing depth", .1, 20)
    p["bearing_capacity_psf"] = _number(p["bearing_capacity_psf"], "Soil bearing capacity (psf)", 100, 100000)
    p["metal_deck_thickness_in"] = _number(p["metal_deck_thickness_in"], "Metal deck thickness", 0, 24)
    p["insulation_depth_in"] = _number(p["insulation_depth_in"], "Insulation depth", 0, 48)
    if p["roof_calculation_mode"] not in ("original", "eave"):
        raise InputValidationError("Roof calculation mode must be original or eave.")
    if p["roof_calculation_mode"] == "original":
        p["roof_height_ft"] = p["clear_height_ft"]
    speed_rows = set()
    for row in _list(p["speed_bay_rows"], "Speed bay rows", MAX_BAYS_PER_AXIS):
        idx = int(row.get("row_index", 1)) - 1 if isinstance(row, dict) else _number(row, "Speed bay row", 0, MAX_BAYS_PER_AXIS - 1, True)
        if 0 <= idx < ny:
            speed_rows.add(idx)
    p["speed_bay_rows"] = sorted(speed_rows)
    p["roof_rise_ft"] = _number(p["roof_rise_ft"], "Roof rise", 0, 100)
    p["joist_seat_depth_in"] = _number(p["joist_seat_depth_in"], "Joist seat depth", 0, 36)
    roof = str(p["roof_type"]).strip().lower().replace("_", " ")
    roof_names = {"flat": "Flat", "gable": "Gable", "double slope": "Gable", "single slope": "Single Slope", "shed": "Single Slope"}
    if roof not in roof_names:
        raise InputValidationError("Roof type must be Flat, Gable, or Single Slope.")
    p["roof_type"] = roof_names[roof]
    p["single_slope_direction"] = "South" if str(p["single_slope_direction"]).lower() == "south" else "North"
    if not isinstance(p["load_inputs_psf"], dict):
        raise InputValidationError("load_inputs_psf must be an object.")
    for key, value in base["load_inputs_psf"].items():
        p["load_inputs_psf"].setdefault(key, value)
    for key, value in p["load_inputs_psf"].items():
        if key.endswith("psf") or key.endswith("psf_manual"):
            p["load_inputs_psf"][key] = _number(value, key, 0, 10000)
    all_bays = {(x, y) for x in range(nx) for y in range(ny)}
    inactive = _bay_indices(p["inactive_bays"], nx, ny, "Inactive bays")
    if "active_bays" in project and "inactive_bays" not in project:
        if not project["active_bays"]:
            raise InputValidationError("At least one bay must remain active.")
        inactive = all_bays - _bay_indices(project["active_bays"], nx, ny, "Active bays")
    p.pop("active_bays", None)
    active = all_bays - inactive
    if not active:
        raise InputValidationError("At least one bay must remain active.")
    p["inactive_bays"] = _serialize_bays(inactive)
    p["collateral_bays"] = _serialize_bays(_bay_indices(p["collateral_bays"], nx, ny, "Collateral bays") & active)
    overrides = {}
    for item in _list(p["joists_per_bay"], "Bay joist spaces"):
        bays = _bay_indices([item], nx, ny, "Bay joist spaces")
        spaces = _number(item.get("joist_spaces", item.get("joist_count", p["joist_spaces"])), "Bay joist spaces", 1, MAX_JOIST_SPACES, True)
        for bay in bays & active:
            overrides[bay] = spaces
    p["joists_per_bay"] = [{"x_bay": x + 1, "y_bay": axis_letter(y), "joist_spaces": count} for (x, y), count in sorted(overrides.items())]
    for layer in _list(p["additional_load_layers"], "Additional load layers", 40):
        if not isinstance(layer, dict):
            raise InputValidationError("Every additional load layer must be an object.")
        layer["psf"] = _number(layer.get("psf", layer.get("load_psf", 0)), "Additional load", 0, 10000)
        layer["bays"] = _serialize_bays(_bay_indices(layer.get("bays", []), nx, ny, "Additional load bays") & active)
    # A named perimeter mode follows footprint edits; imported explicit segments
    # stay explicit until the user chooses a global mode.
    if p.get("perimeter_support") == "walls":
        p["load_bearing_wall_segments"] = [
            {"orientation": orient, "line_index": line + 1, "segment_index": segment + 1}
            for orient, line, segment in sorted(engine._active_boundary_segment_set(active, p["x_spans_ft"], p["y_spans_ft"]))
        ]
    elif p.get("perimeter_support") == "steel":
        p["load_bearing_wall_segments"] = []
    for edge in _list(p["load_bearing_wall_segments"], "Load-bearing wall segments", 2400):
        if not isinstance(edge, dict):
            raise InputValidationError("Every wall segment must be an object.")
        if str(edge.get("orientation", "")).upper() not in ("H", "V"):
            raise InputValidationError("Wall orientation must be H or V.")
        _number(edge.get("line_index"), "Wall line", 1, MAX_BAYS_PER_AXIS + 1, True)
        _number(edge.get("segment_index"), "Wall segment", 1, MAX_BAYS_PER_AXIS, True)
    zones = _list(p["mezzanines"], "Mezzanines", 20)
    zone_ids = set()
    estimated_members = sum(overrides.get(bay, p["joist_spaces"]) + 4 for bay in active)
    for index, zone in enumerate(zones):
        if not isinstance(zone, dict):
            raise InputValidationError("Each mezzanine must be an object.")
        zone["id"] = str(zone.get("id") or f"MZ{index + 1:03d}")[:100]
        if zone["id"] in zone_ids:
            raise InputValidationError("Mezzanine IDs must be unique.")
        zone_ids.add(zone["id"])
        zone["elevation_ft"] = _number(zone.get("elevation_ft", 15), "Mezzanine elevation", 1, p["roof_height_ft"] - .5)
        for key in ("dead_load_psf", "live_load_psf"):
            zone[key] = _number(zone.get(key, 0), f"Mezzanine {key}", 0, 10000)
        zone["joist_spaces_per_bay"] = _number(zone.get("joist_spaces_per_bay", 7), "Mezzanine joist spaces", 1, MAX_JOIST_SPACES, True)
        for key in ("internal_x_spacings_ft", "internal_y_spacings_ft"):
            vals = zone.get(key, [])
            if isinstance(vals, str):
                vals = [v.strip() for v in vals.split(",") if v.strip()]
            zone[key] = [_number(v, key, 1, 200) for v in _list(vals, key, MAX_BAYS_PER_AXIS)]
        for key in ("x_start_ft", "x_end_ft", "y_start_ft", "y_end_ft"):
            if key in zone:
                zone[key] = _number(zone[key], key, 0, 4800)
        panel_overrides = zone.get("joist_spaces_overrides") or {}
        if not isinstance(panel_overrides, dict):
            raise InputValidationError("Mezzanine joist_spaces_overrides must be an object.")
        for value in panel_overrides.values():
            _number(value, "Mezzanine panel joist spaces", 1, MAX_JOIST_SPACES, True)
        if p["mezzanine_enabled"] and zone.get("enabled", True):
            # Segment fitting adds at most one remainder panel per axis.
            panels = (len(zone["internal_x_spacings_ft"]) + 1) * (len(zone["internal_y_spacings_ft"]) + 1)
            max_spaces = max([zone["joist_spaces_per_bay"]] + [int(v) for v in panel_overrides.values()])
            estimated_members += panels * (max_spaces + 4)
    for key in ("assignments", "member_overrides", "group_overrides", "default_sections"):
        if not isinstance(p[key], dict):
            raise InputValidationError(f"{key} must be an object.")
    for value in p["assignments"].values():
        if not isinstance(value, dict):
            raise InputValidationError("Each assignments category must be an object.")
    if estimated_members > MAX_MEMBERS:
        raise InputValidationError("This layout is too large for the interactive model. Reduce bays or joist spaces.")
    layout = copy.deepcopy(p)
    layout["active_bays"] = _serialize_bays(active)
    layout["joists_per_bay"] = [{"x_bay": x + 1, "y_bay": axis_letter(y), "joist_count": overrides.get((x, y), p["joist_spaces"])} for x, y in sorted(active)]
    return p, layout


def _catalog_number(value):
    if value is None:
        return None
    try:
        value = float(str(value).replace(",", "").strip())
        return value if math.isfinite(value) else None
    except ValueError:
        return None


def _section(kind, designation, depth, weight, source, **extra):
    item = {"id": designation, "designation": designation, "depth_in": depth,
            "weight_plf": weight, "source": source, "description": DESCRIPTIONS[kind],
            "profile": "hss_square" if kind == "column" else "truss", "type": kind,
            "grade": "Unspecified", "check_status": "unchecked"}
    if kind == "column":
        dims = re.search(r"HSS\s*(\d+(?:\.\d+)?)\s*x\s*(\d+(?:\.\d+)?)", designation, re.I)
        if dims:
            item["width_in"] = float(dims.group(2))
    if re.match(r"^W\d", designation, re.I):
        item["profile"] = "wide_flange"
        item["description"] = "Wide-flange steel section. Two flanges connected by a solid web form an I-shaped profile. Depth and weight are selection inputs; steel grade, capacity and connection details are unchecked."
    item.update(extra)
    return item


@lru_cache(maxsize=1)
def _catalog_data():
    """Read the same workbook layouts as grid_gui, without importing Tkinter."""
    output = {"joist": [], "girder": [], "column": []}
    notices = []
    try:
        from openpyxl import load_workbook
    except ImportError:
        return output, ["Install openpyxl to load the supplied steel catalogs. Custom section entries remain available."]
    root = Path(__file__).resolve().parent
    for filename in ("Joist Table 2.xlsx", "LH Joist Table.xlsx", "Column Table.xlsx", "Expanded Vulcraft Joist Girder Catalog.xlsx"):
        path = root / filename
        if not path.is_file():
            notices.append(f"Catalog unavailable: {filename}.")
            continue
        wb = None
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            rows = list(wb.active.iter_rows(values_only=True))
            if "Joist Table" in filename:
                header = next(i for i, row in enumerate(rows[:100]) if "span" in str(row[0]).lower() and any(re.fullmatch(r"\d{2,3}(?:K|LH)\d{1,2}", str(c), re.I) for c in row))
                for col, raw in enumerate(rows[header]):
                    designation = str(raw or "").strip()
                    if not re.fullmatch(r"\d{2,3}(?:K|LH)\d{1,2}", designation, re.I):
                        continue
                    depth, weight = _catalog_number(rows[header + 1][col]), _catalog_number(rows[header + 2][col])
                    if not depth or not weight:
                        continue
                    capacities = []
                    for row in rows[header + 3:]:
                        span, capacity = _catalog_number(row[0]), _catalog_number(row[col])
                        if span and capacity and 10 <= span <= 150 and capacity > 0:
                            capacities.append({"span_ft": span, "capacity_plf": capacity})
                    if capacities:
                        output["joist"].append(_section("joist", designation, depth, weight, filename, capacities=capacities))
            elif filename == "Column Table.xlsx":
                for col, raw in enumerate(rows[1]):
                    designation = str(raw or "").strip()
                    if not re.fullmatch(r"\d+(?:\.\d+)?x\d+(?:\.\d+)?x\d+/\d+", designation) or col < 1:
                        continue
                    if str(rows[14][col - 1]).strip().upper() != "ASD":
                        continue
                    weight = _catalog_number(rows[2][col])
                    if not weight:
                        continue
                    capacities = []
                    for row in rows[16:]:
                        kl, capacity = _catalog_number(row[0]), _catalog_number(row[col - 1])
                        if kl and capacity and kl >= 1 and capacity > 0:
                            capacities.append({"kl_ft": kl, "asd_capacity_kips": capacity})
                    output["column"].append(_section("column", "HSS" + designation, float(designation.split("x")[0]), weight, filename, capacities=capacities))
            else:
                header, start, loads = None, None, []
                for i, row in enumerate(rows[:40]):
                    for j in range(4, len(row)):
                        trial = []
                        for raw in row[j:]:
                            value = _catalog_number(raw)
                            if value is None or not 1 <= value <= 100:
                                break
                            trial.append(value)
                        if len(trial) >= 4:
                            header, start, loads = i, j, trial
                            break
                    if header is not None:
                        break
                if header is None:
                    raise ValueError("No panel-load headers found")
                span, spaces = None, ""
                for row in rows[header + 1:]:
                    candidate_span = _catalog_number(row[1])
                    if candidate_span and 10 <= candidate_span <= 400:
                        span = candidate_span
                    if isinstance(row[2], str) and row[2].strip():
                        spaces = row[2].strip()
                    depth = _catalog_number(row[3])
                    match = re.search(r"(\d+)\s*N", spaces, re.I)
                    if not span or not depth or depth <= 0 or not match:
                        continue
                    n = int(match.group(1))
                    for offset, panel_load in enumerate(loads):
                        weight = _catalog_number(row[start + offset])
                        if weight is None or weight <= 0:
                            continue
                        designation = f"{depth:g}G {n}N {panel_load:g}K"
                        output["girder"].append(_section("girder", designation, depth, weight, filename,
                            id=f"{designation}@{span:g}ft", span_ft=span, joist_n=n, panel_load_kips=panel_load))
        except Exception as exc:
            notices.append(f"Could not read {filename}: {exc}")
        finally:
            if wb is not None:
                wb.close()
    for kind in output:
        unique = {item["id"]: item for item in output[kind]}
        output[kind] = sorted(unique.values(), key=lambda s: (s["depth_in"], s["weight_plf"], s["id"]))
    return output, notices


def catalogs():
    """Return local catalog sections; capacities are catalog data, not checks."""
    return copy.deepcopy(_catalog_data()[0])


def _lines(spans, origin=0):
    result = [float(origin)]
    for span in spans:
        result.append(result[-1] + float(span))
    return result


def _roof_heights(p, results):
    ys = _lines(p["y_spans_ft"])
    if p.get("roof_line_elevations_ft") is not None:
        supplied = p["roof_line_elevations_ft"]
        if not isinstance(supplied, list) or len(supplied) != len(ys):
            raise InputValidationError("roof_line_elevations_ft must contain one elevation per Y grid line.")
        return [_number(v, "Roof line elevation", 4, 300) for v in supplied]
    if p.get("legacy_roof_profile") and p["roof_type"] != "Flat":
        # Reproduce grid_gui's line-profile method (1/4 in/ft main slope;
        # 1/2 in/ft broken-clear-height speed bays), including its anchor.
        speed = set()
        for value in p.get("speed_bay_rows", []):
            idx = int(value.get("row_index", 1)) - 1 if isinstance(value, dict) else int(value)
            if 0 <= idx < len(ys) - 1:
                speed.add(idx)
        transitions = [i for i in range(1, len(ys) - 1) if ((i - 1) in speed) != (i in speed)]
        direction = p["single_slope_direction"]
        broken = bool((p.get("inputs") or {}).get("break_clear_height", False))
        ridge = len(ys) // 2
        if broken and transitions:
            anchor = (max(transitions) if direction == "North" else min(transitions)) if p["roof_type"] == "Single Slope" else min(transitions, key=lambda i: abs(i - ridge))
        else:
            anchor = len(ys) - 1 if direction == "North" else 0
        row = max(0, anchor - 1)
        if broken and transitions:
            north, south = anchor - 1, anchor
            if south < len(ys) - 1 and south not in speed and (north < 0 or north in speed):
                row = south
            elif not (north >= 0 and north not in speed):
                row = south if direction == "North" else north
        row = min(len(ys) - 2, max(0, row))
        assignments = p["assignments"].get("joists", {})
        depths = [float(assignments.get(f"{bay['required_capacity_plf']:.2f}|{bay['bay_length_ft']:.2f}", {}).get("depth_in", 0) or 0)
                  for bay in results["joists"]["joist_bay_calculations"] if axis_index(bay["y_row_label"]) == row]
        height = p["roof_height_ft"] + max(depths or [0]) / 12
        slopes = []
        for i in range(len(ys) - 1):
            sign = (1 if i < ridge else -1) if p["roof_type"] == "Gable" else (-1 if direction == "North" else 1)
            if broken and i in speed:
                sides = []
                if i > 0 and i - 1 not in speed:
                    sides.append("South")
                if i < len(ys) - 2 and i + 1 not in speed:
                    sides.append("North")
                side = sides[0] if len(sides) == 1 else (("South" if direction == "North" else "North") if len(sides) > 1 else ("North" if i < (len(ys) - 1) / 2 else "South"))
                slopes.append((1 if side == "North" else -1) * .5 / 12)
            else:
                slopes.append(sign * .25 / 12)
        heights = [height] * len(ys)
        for i in range(anchor - 1, -1, -1):
            heights[i] = heights[i + 1] - slopes[i] * p["y_spans_ft"][i]
        for i in range(anchor, len(ys) - 1):
            heights[i + 1] = heights[i] + slopes[i] * p["y_spans_ft"][i]
        if min(heights) <= 0:
            raise InputValidationError("Imported roof profile falls below the floor. Increase clear height.")
        return heights
    height, rise = p["roof_height_ft"], p["roof_rise_ft"]
    if p["roof_type"] == "Flat":
        return [height] * len(ys)
    if p["roof_type"] == "Single Slope":
        return [height + rise * ((1 - y / ys[-1]) if p["single_slope_direction"] == "North" else y / ys[-1]) for y in ys]
    # Ridge follows a grid line, so no joist is secretly split into two members.
    if len(ys) < 3:
        raise InputValidationError("A gable roof needs at least two Y bays so the ridge has a support line.")
    ridge = ys[len(ys) // 2]
    return [height + rise * (y / ridge if y <= ridge else (ys[-1] - y) / (ys[-1] - ridge)) for y in ys]


def _zone_nodes(p, zone):
    source = next(item for item in p["mezzanines"] if str(item["id"]) == zone["zone_id"])
    xs, ys = _lines(p["x_spans_ft"]), _lines(p["y_spans_ft"])
    x0 = min(source["x_start_ft"], source["x_end_ft"]) if "x_start_ft" in source else xs[zone["x_start_line_index"] - 1]
    y0 = min(source["y_start_ft"], source["y_end_ft"]) if "y_start_ft" in source else ys[zone["y_start_line_index"] - 1]
    return _lines(zone["internal_x_spacings_ft"], max(0, min(x0, xs[-1]))), _lines(zone["internal_y_spacings_ft"], max(0, min(y0, ys[-1])))


def _raw_members(p, layout, results, heights):
    xs, ys = _lines(p["x_spans_ft"]), _lines(p["y_spans_ft"])
    counts = engine._joist_count_map(layout["joists_per_bay"])
    members = []

    def add(kind, member_id, engine_id, row, start, end, group, level="roof"):
        members.append({"id": member_id, "engine_id": engine_id, "type": kind, "level": level,
                        "start": [round(v, 6) for v in start], "end": [round(v, 6) for v in end],
                        "group_id": f"{level}:{kind}:{group}", "assignment_group_id": group,
                        "demand": copy.deepcopy(row), "label": engine_id,
                        "bay": row.get("bay", row.get("zone_name", row.get("line_label", "")))})
    for row in results["joists"]["joist_members"]:
        x, y = row["x_bay_index"] - 1, axis_index(row["y_row_label"])
        index = row["member_index_in_bay"] - 1
        spaces = counts[(x, y)]
        pos = xs[x] + index * p["x_spans_ft"][x] / spaces
        if index == 0 or index == spaces:
            identity = f"J-{axis_letter(y)}-L{x + (1 if index == spaces else 0) + 1}"
        else:
            identity = f"J-{axis_letter(y)}{x + 1}-{index}of{spaces}"
        group = f"{row['required_capacity_plf']:.2f}|{row['span_ft']:.2f}"
        add("joist", identity, row["id"], row, [pos, ys[y], heights[y]], [pos, ys[y + 1], heights[y + 1]], group)
    for row in results["girders"]["girder_calculations"]:
        x, y = row["x_bay_index"] - 1, axis_index(row["line_label"])
        group = f"{row['required_capacity_lbs']:.2f}|{row['bay_width_ft']:.2f}|{row['required_joist_count_n']}"
        # Native GUI groups include allowed depth. Retain exact saved assignment
        # when its demand/span/panel prefix identifies one existing group.
        matches = [g for g in {**p.get("auto_assignments", {}).get("girders", {}), **p["assignments"].get("girders", {})} if str(g).startswith(group + "|")]
        if str(y) in p.get("girder_depth_limits_in", {}):
            group += f"|{p['girder_depth_limits_in'][str(y)]:.2f}"
        elif p.get("legacy_roof_profile"):
            clear = float((p.get("inputs") or {}).get("clear_height_ft", p["roof_height_ft"]))
            exact = group + f"|{(heights[y] - clear) * 12:.2f}"
            group = exact if exact in matches else (matches[0] if len(matches) == 1 else group)
        elif len(matches) == 1:
            group = matches[0]
        z = heights[y] - p["joist_seat_depth_in"] / 12
        add("girder", row["girder_id"], row["girder_id"], row, [xs[x], ys[y], z], [xs[x + 1], ys[y], z], group)
    for row in results["columns"]["column_calculations"]:
        x, y = row["x_line_index"], row["y_line_index"]
        add("column", row["column_id"], row["column_id"], row, [xs[x], ys[y], 0], [xs[x], ys[y], heights[y] - p["joist_seat_depth_in"] / 12], f"{row['required_capacity_kips']:.2f}")
    mezz = results["mezzanine"]
    zones = {z["zone_id"]: (z, *_zone_nodes(p, z)) for z in mezz["mezzanine_zones"]}
    panels = {r["panel_id"]: r for r in mezz["mezzanine_panel_calculations"]}
    for row in mezz["mezzanine_joist_calculations"]:
        zone, zx, zy = zones[row["zone_id"]]
        panel = panels[row["panel_id"]]
        ix, iy = panel["panel_x_index"] - 1, panel["panel_y_index"] - 1
        spaces = row["joist_spaces_per_bay"]
        vertical = row["joist_direction"] == "vertical"
        edge_a = (ix == 0 and not zone["west_lb_wall"]) if vertical else (iy == 0 and not zone["north_lb_wall"])
        edge_b = (ix < len(zx) - 2 or not zone["east_lb_wall"]) if vertical else (iy < len(zy) - 2 or not zone["south_lb_wall"])
        indices = ([0] if edge_a else []) + list(range(1, spaces)) + ([spaces] if edge_b else [])
        if len(indices) != row["joist_count"]:
            raise RuntimeError("Mezzanine geometry does not match engine joist count.")
        for index in indices:
            z = zone["elevation_ft"]
            if vertical:
                x = zx[ix] + (zx[ix + 1] - zx[ix]) * index / spaces
                start, end = [x, zy[iy], z], [x, zy[iy + 1], z]
            else:
                y = zy[iy] + (zy[iy + 1] - zy[iy]) * index / spaces
                start, end = [zx[ix], y, z], [zx[ix + 1], y, z]
            identity = f"{row['panel_id']}-J{index}of{spaces}"
            group = row.get("group_id", f"{row['required_capacity_plf']:.2f}|{row['joist_span_ft']:.2f}")
            add("joist", identity, identity, row, start, end, group, "mezzanine")
    for row in mezz["mezzanine_girder_calculations"]:
        zone, zx, zy = zones[row["zone_id"]]
        line, segment = [int(v) - 1 for v in row["girder_id"].rsplit("-", 2)[1:]]
        z = zone["elevation_ft"] - p["joist_seat_depth_in"] / 12
        if row["joist_direction"] == "vertical":
            start, end = [zx[segment], zy[line], z], [zx[segment + 1], zy[line], z]
        else:
            start, end = [zx[line], zy[segment], z], [zx[line], zy[segment + 1], z]
        add("girder", row["girder_id"], row["girder_id"], row, start, end, row.get("group_id", row["girder_id"]), "mezzanine")
    for row in mezz["mezzanine_column_calculations"]:
        x, y, z = row["x_ft"], row["y_ft"], row["column_height_ft"]
        add("column", row["column_id"], row["column_id"], row, [x, y, 0], [x, y, z - p["joist_seat_depth_in"] / 12], row.get("group_id", row["column_id"]), "mezzanine")
    return members


def _selected_section(member, p, catalog_index):
    kind, level = member["type"], member["level"]
    assignment_key = ("mezz_" if level == "mezzanine" else "") + {"joist": "joists", "girder": "girders", "column": "columns"}[kind]
    selected = p["member_overrides"].get(member["id"])
    source = "Member selection"
    if selected is None:
        selected = p["group_overrides"].get(member["group_id"])
        source = "Group selection"
    if selected is None:
        selected = p["assignments"].get(assignment_key, {}).get(member["assignment_group_id"])
        source = "Imported assignment"
    if selected is None:
        selected = p["default_sections"].get(kind)
        source = "Project section selection"
    if selected is None:
        selected = p.get("auto_assignments", {}).get(assignment_key, {}).get(member["assignment_group_id"])
        source = "Original automatic catalog selection"
    if selected is None or selected == "" or selected == {}:
        return _section(kind, "Unassigned", {"joist": 30, "girder": 48, "column": 12}[kind], None,
                        "Schematic placeholder", id="unassigned", assigned=False,
                        geometry_status="Illustrative profile; choose a section to set its depth.")
    if isinstance(selected, str):
        selected = {"section_id": selected}
    if not isinstance(selected, dict):
        raise InputValidationError("A section override must be a section ID or an object.")
    section_id = selected.get("section_id", selected.get("id"))
    entry = catalog_index[kind].get(str(section_id)) if section_id else None
    if section_id and entry is None:
        raise InputValidationError(f"Unknown {kind} catalog section: {section_id}.")
    if entry is None and selected.get("designation"):
        entry = catalog_index[kind].get(selected["designation"])
    result = copy.deepcopy(entry or {})
    result.update({k: v for k, v in selected.items() if k not in ("section_id", "check_status", "assigned")})
    depth = _number(result.get("depth_in"), "Section depth", .1, 180)
    raw_weight = result.get("weight_plf")
    weight = _number(raw_weight, "Section weight", .001, 10000) if raw_weight is not None else None
    designation = str(result.get("designation") or (f"{depth:g}G" if kind == "girder" else f"Custom {kind}"))[:150]
    section = _section(kind, designation, depth, weight, result.get("source", source))
    section.update(result)
    section.update(depth_in=depth, weight_plf=weight, designation=designation, assigned=True,
                   selection_source=source, check_status="unchecked", geometry_status="Schematic geometry; connection and web details are illustrative.")
    if "width_in" in section:
        section["width_in"] = _number(section["width_in"], "Section width", .1, 180)
    if section.get("profile") not in ("truss", "wide_flange", "hss_round", "hss_square"):
        section["profile"] = "hss_square" if kind == "column" else "truss"
    return section


def _project_catalog(catalog, members):
    """Keep the girder chooser useful: show one catalog row per depth and demand.

    Rows use the next available span and panel load at the required panel count.
    This is just a catalog browsing filter, with every member still unchecked.
    """
    keys = set()
    for member in members:
        if member["type"] == "girder":
            row = member["demand"]
            keys.add((row.get("bay_width_ft", row.get("span_ft", 0)), row.get("required_joist_count_n", row.get("required_joist_n", 0)), row.get("required_capacity_lbs", 0) / 1000))
    chosen = {}
    for span, n, load in keys:
        by_depth = {}
        for item in catalog["girder"]:
            if item["joist_n"] != n or item["span_ft"] < span or item["panel_load_kips"] < load:
                continue
            current = by_depth.get(item["depth_in"])
            if current is None or (item["span_ft"], item["panel_load_kips"]) < (current["span_ft"], current["panel_load_kips"]):
                by_depth[item["depth_in"]] = item
        chosen.update({item["id"]: item for item in by_depth.values()})
    for member in members:
        if member["type"] == "girder" and member.get("section", {}).get("assigned"):
            item = next((s for s in catalog["girder"] if s["id"] == member["section"]["id"]), None)
            if item:
                chosen[item["id"]] = item
    return {"joist": catalog["joist"], "column": catalog["column"],
            "girder": sorted(chosen.values(), key=lambda v: (v["depth_in"], v["weight_plf"]))}


def calculate_project(project):
    """Normalize a new/native saved project, run the engine and build its model."""
    try:
        p, layout = normalize_project(project)
        results = {"joists": engine.calculate_joist_takeoff(layout),
                   "girders": engine.calculate_girder_takeoff(layout),
                   "columns": engine.calculate_column_takeoff(layout),
                   "mezzanine": engine.calculate_mezzanine_takeoff(layout)}
        catalog, catalog_notices = _catalog_data()
        index = {kind: {s["id"]: s for s in rows} for kind, rows in catalog.items()}
        from takeoff_workflow import Workflow
        workflow = Workflow(p, results)
        workflow.select_joists()
        p["auto_assignments"] = workflow.auto
        # Explicit member/group selections govern roof clearances too.
        preliminary = _raw_members(p, layout, results, [p["roof_height_ft"]] * (len(p["y_spans_ft"]) + 1))
        effective_joists = {}
        for member in preliminary:
            if member["type"] == "joist" and member["level"] == "roof":
                section = _selected_section(member, p, index)
                if section["assigned"]:
                    gid = member["assignment_group_id"]
                    if section["depth_in"] >= effective_joists.get(gid, {}).get("depth_in", 0):
                        effective_joists[gid] = section
        workflow.joist_selection_by_group = effective_joists
        if p["roof_calculation_mode"] == "original":
            if p["roof_type"] == "Flat":
                depth = max((s["depth_in"] for s in effective_joists.values()), default=0)
                heights = [p["clear_height_ft"] + depth / 12] * (len(p["y_spans_ft"]) + 1)
            else:
                heights = workflow._build_roof_profile_data()["line_toj_elevations_ft"]
        else:
            heights = _roof_heights(p, results)
        workflow.set_profile(heights)
        workflow.select_remaining()
        p["girder_depth_limits_in"] = {str(k): v for k, v in workflow._get_allowed_girder_depth_by_line_in().items()}
        extra_takeoffs = workflow.concrete_and_walls()
        members = _raw_members(p, layout, results, heights)
        if len(members) > MAX_MEMBERS:
            raise InputValidationError(f"Model exceeds {MAX_MEMBERS:,} members. Reduce bays, mezzanine panels, or joist spaces.")
        total_length = total_weight = selected_weight = estimated_weight = 0.0
        selected_count = 0
        weight_known_count = 0
        for member in members:
            section = _selected_section(member, p, index)
            member["section"] = section
            member["description"] = section["description"]
            member["check_status"] = "unchecked"
            member["length_ft"] = round(math.dist(member["start"], member["end"]), 4)
            row = member["demand"]
            member["weight_length_ft"] = (member["length_ft"] if member["type"] == "column" else
                float(row.get("span_ft", row.get("joist_span_ft", row.get("bay_width_ft", member["length_ft"])))))
            total_length += member["length_ft"]
            member["weight_lbs"] = None
            if section["assigned"]:
                selected_count += 1
                if section["weight_plf"] is not None:
                    member["weight_lbs"] = round(member["weight_length_ft"] * section["weight_plf"], 2)
                    total_weight += member["weight_lbs"]
                    weight_known_count += 1
                    # Imported interpolation/extrapolation is retained as an estimate.
                    if section.get("tier", section.get("selection_tier", "exact")) not in ("exact", ""):
                        estimated_weight += member["weight_lbs"]
                    else:
                        selected_weight += member["weight_lbs"]
        xs, ys = _lines(p["x_spans_ft"]), _lines(p["y_spans_ft"])
        active = engine._active_bay_set(layout, p["x_spans_ft"], p["y_spans_ft"])
        collateral = engine._bay_set(p["collateral_bays"])
        boundary = engine._active_boundary_segment_set(active, p["x_spans_ft"], p["y_spans_ft"])
        wall_segments = engine._boundary_load_bearing_set(layout, p["x_spans_ft"], p["y_spans_ft"], active) & boundary
        walls = []
        for orient, line, segment in sorted(wall_segments):
            if orient == "H":
                start, end = [xs[segment], ys[line], heights[line]], [xs[segment + 1], ys[line], heights[line]]
            else:
                start, end = [xs[line], ys[segment], heights[segment]], [xs[line], ys[segment + 1], heights[segment + 1]]
            walls.append({"id": f"W-{orient}-{line + 1}-{segment + 1}", "start": start, "end": end})
        bays = [{"id": f"{axis_letter(y)}{x + 1}", "x_bay": x + 1, "y_bay": axis_letter(y),
                 "x_index": x, "y_index": y, "x_ft": xs[x], "y_ft": ys[y],
                 "width_ft": p["x_spans_ft"][x], "length_ft": p["y_spans_ft"][y],
                 "active": (x, y) in active, "collateral": (x, y) in collateral,
                 "corners": [[xs[x], ys[y], heights[y]], [xs[x + 1], ys[y], heights[y]],
                             [xs[x + 1], ys[y + 1], heights[y + 1]], [xs[x], ys[y + 1], heights[y + 1]]]}
                for y in range(len(ys) - 1) for x in range(len(xs) - 1)]
        notices = list(catalog_notices)
        notices.extend(workflow.notes)
        notices.append("Takeoff demands use the existing calculation_engine.py. Section adequacy, stability, deflection, connections and steel grade are unchecked.")
        if p["roof_calculation_mode"] == "original":
            notices.append("Roof follows the original clear-height, assigned-joist-depth and speed-bay slope rules. Girders use the original depth limits by roof line. Use Explicit eave height & rise to set elevations directly.")
        else:
            notices.append("Roof elevation is top of joist at the eave; girder support is lowered by the joist seat depth. Roof slope changes member lengths; engine demands use plan spans.")
        if p["mezzanine_enabled"]:
            notices.append("Mezzanine members follow the engine's support and shared-edge rules. Reused main columns are not counted twice; combined roof + mezzanine demand on main columns is not checked by this engine.")
        summary = {"member_count": len(members), "bay_count": len(active),
                   "floor_area_sf": round(sum(p["x_spans_ft"][x] * p["y_spans_ft"][y] for x, y in active), 3),
                   "mezzanine_area_sf": results["mezzanine"]["summary"]["total_mezzanine_area_sf"],
                   "total_length_ft": round(total_length, 2), "selected_weight_lbs": round(selected_weight, 2),
                   "estimated_weight_lbs": round(estimated_weight, 2), "total_weight_lbs": round(total_weight, 2),
                   "weight_tons": round(total_weight / 2000, 3), "selected_count": selected_count,
                   "unassigned_count": len(members) - selected_count,
                   "unknown_weight_count": len(members) - weight_known_count,
                   "unchecked_count": len(members), "check_status": "unchecked",
                   "weight_status": "Partial selected-member takeoff" if weight_known_count < len(members) else "Selected-member takeoff; connections excluded"}
        for kind in ("joist", "girder", "column"):
            summary[kind + "_count"] = sum(m["type"] == kind for m in members)
        model = {"project": p, "members": members, "results": results, "bays": bays, "walls": walls,
                 "takeoffs": extra_takeoffs, "roof_profile": workflow.profile,
                 "selection_groups": workflow.groups,
                 "grid": {"x_lines_ft": xs, "y_lines_ft": ys, "x_lines": xs, "y_lines": ys,
                          "width_ft": xs[-1], "length_ft": ys[-1], "roof_heights_ft": heights,
                          "height_ft": max(heights)},
                 "summary": summary, "catalog": _project_catalog(catalog, members), "notices": notices,
                 "units": {"length": "ft", "section_depth": "in", "weight": "lb", "area_load": "psf"}}
        _finite_tree(model)
        return model
    except (KeyError, TypeError, OverflowError, AttributeError) as exc:
        raise InputValidationError(f"Invalid project field: {exc}") from exc
